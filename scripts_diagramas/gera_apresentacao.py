# -*- coding: utf-8 -*-
"""Gera a apresentacao do trabalho (Apresentacao_Trabalho_BD.pptx)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(BASE, '..', 'documentacao')

NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
LIGHT = RGBColor(0xF2, 0xF4, 0xFA)
GREEN = RGBColor(0x1E, 0x6B, 0x2C)
RED = RGBColor(0x8F, 0x23, 0x1C)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def tx(s, x, y, w, h, lines, size=15, color=DARK, bold=False, font='Calibri',
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=6, line_sp=None):
    """lines: str ou lista de (texto, dict_overrides)"""
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [(lines, {})]
    for i, (texto, ov) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ov.get('align', align)
        p.space_after = Pt(ov.get('sp_after', sp_after))
        if line_sp:
            p.line_spacing = line_sp
        r = p.add_run()
        r.text = texto
        r.font.size = Pt(ov.get('size', size))
        r.font.bold = ov.get('bold', bold)
        r.font.name = ov.get('font', font)
        r.font.color.rgb = ov.get('color', color)
    return box


def titulo(s, texto, sub=None):
    tx(s, Inches(0.6), Inches(0.42), Inches(12.1), Inches(0.8),
       texto, size=31, bold=True, color=NAVY, font='Georgia')
    if sub:
        tx(s, Inches(0.6), Inches(1.12), Inches(12.1), Inches(0.4),
           sub, size=14, color=GRAY)


def card(s, x, y, w, h, fill=LIGHT, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    # raio menor
    sh.adjustments[0] = 0.06
    return sh


def codebox(s, x, y, w, h, codigo_, size=11):
    c = card(s, x, y, w, h, fill=RGBColor(0xF7, 0xF8, 0xFB), line=ICE)
    tf = c.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = tf.margin_bottom = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, l in enumerate(codigo_.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = l if l else ' '
        r.font.size = Pt(size)
        r.font.name = 'Consolas'
        r.font.color.rgb = DARK
    return c


def numcirc(s, x, y, n, d=Inches(0.42), fill=NAVY):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return c


def stat(s, x, y, w, numero, rotulo, cor=NAVY, num_size=54):
    tx(s, x, y, w, Inches(0.9), numero, size=num_size, bold=True, color=cor,
       font='Georgia', align=PP_ALIGN.CENTER)
    tx(s, x, y + Inches(0.95), w, Inches(0.6), rotulo, size=13, color=GRAY,
       align=PP_ALIGN.CENTER)


# ================= SLIDE 1 - CAPA =================
s = slide(NAVY)
tx(s, Inches(1), Inches(2.0), Inches(11.3), Inches(1.6),
   'Sistema de Gestão de Horários de Disciplinas',
   size=44, bold=True, color=WHITE, font='Georgia', align=PP_ALIGN.CENTER)
tx(s, Inches(1), Inches(3.55), Inches(11.3), Inches(0.6),
   'Modelagem e implementação de banco de dados para alocação de horários do curso de ADS',
   size=18, color=ICE, align=PP_ALIGN.CENTER)
tx(s, Inches(1), Inches(5.6), Inches(11.3), Inches(1.2),
   [('Trabalho Final — Banco de Dados — 1/2026', {'sp_after': 4}),
    ('PUC Minas · ICEI · Tecnologia em Análise e Desenvolvimento de Sistemas', {'size': 13})],
   size=15, color=ICE, align=PP_ALIGN.CENTER)

# ================= SLIDE 2 - O PROBLEMA =================
s = slide()
titulo(s, 'O problema: montar a grade de horários do ADS',
       'Um problema clássico de alocação de recursos com restrições (timetabling)')
regras = [
    ('1 professor por disciplina', 'cada disciplina é ministrada por exatamente um professor habilitado'),
    ('Sem choque de período', 'disciplinas do mesmo período não podem ocupar o mesmo horário'),
    ('Professor sem ubiquidade', 'um professor não pode estar em duas salas ao mesmo tempo'),
    ('2 horários por dia', '5 dias × 2 slots = apenas 10 horários disponíveis na semana'),
]
yy = Inches(1.85)
for i, (t, d) in enumerate(regras):
    numcirc(s, Inches(0.7), yy, i + 1)
    tx(s, Inches(1.35), yy - Inches(0.04), Inches(6.2), Inches(0.95),
       [(t, {'bold': True, 'size': 16, 'color': NAVY, 'sp_after': 2}),
        (d, {'size': 13.5, 'color': GRAY})])
    yy += Inches(1.05)
card(s, Inches(8.1), Inches(1.9), Inches(4.5), Inches(3.9), fill=LIGHT)
tx(s, Inches(8.1), Inches(2.25), Inches(4.5), Inches(0.5),
   'OBJETIVO', size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
tx(s, Inches(8.35), Inches(2.8), Inches(4.0), Inches(2.6),
   'Maximizar o número de disciplinas em paralelo — compactar a grade '
   'aproveitando ao máximo cada horário, sem violar nenhuma restrição.',
   size=16, color=DARK, align=PP_ALIGN.CENTER)
tx(s, Inches(8.35), Inches(4.6), Inches(4.0), Inches(1.0),
   '20 disciplinas · 8 professores · 6 salas · 10 horários',
   size=13, color=GRAY, align=PP_ALIGN.CENTER)

# ================= SLIDE 3 - VISAO GERAL =================
s = slide()
titulo(s, 'Linhas gerais da solução')
etapas = [
    ('Modelagem conceitual', 'DER com 12 entidades + diagrama de classes'),
    ('Modelo lógico', '13 tabelas normalizadas (3FN)'),
    ('Modelo físico', 'DDL T-SQL: PKs, FKs, UNIQUE, CHECK e 2 triggers'),
    ('Consultas SQL', '12 consultas + 2 visões (≥ 3 tabelas cada)'),
    ('Algoritmo + aplicação', 'coloração de grafos + Flask conectado ao SQL Server'),
]
xx = Inches(0.55)
w = Inches(2.35)
for i, (t, d) in enumerate(etapas):
    c = card(s, xx, Inches(2.3), w, Inches(2.5), fill=LIGHT if i % 2 == 0 else ICE)
    numcirc(s, xx + w / 2 - Inches(0.21), Inches(2.62), i + 1)
    tx(s, xx + Inches(0.12), Inches(3.3), w - Inches(0.24), Inches(0.8),
       t, size=14.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    tx(s, xx + Inches(0.14), Inches(4.0), w - Inches(0.28), Inches(0.9),
       d, size=11.5, color=GRAY, align=PP_ALIGN.CENTER)
    xx += w + Inches(0.16)
tx(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(0.6),
   'SGBD: Microsoft SQL Server   ·   Aplicação: Python + Flask + pyodbc   ·   Interface: HTML/Jinja2',
   size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ================= SLIDE 4 - DER =================
s = slide()
titulo(s, 'Modelagem conceitual — DER')
s.shapes.add_picture(os.path.join(DOC, 'DER.png'), Inches(0.4), Inches(2.1), width=Inches(12.5))
tx(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.3),
   [('ALOCACAO é o núcleo do modelo: entidade associativa que liga DISCIPLINA, PROFESSOR, '
     'HORARIO, SALA e SEMESTRE_LETIVO — representa uma "aula publicada" na grade.',
     {'sp_after': 4}),
    ('USUARIO generaliza PROFESSOR e ALUNO (especializações 1:1) · '
     'HORARIO pertence a um DIA_SEMANA com slot 1 ou 2 (regra dos 2 horários/dia).',
     {'color': GRAY, 'size': 13})], size=14.5)

# ================= SLIDE 5 - DECISOES DE MODELAGEM =================
s = slide()
titulo(s, 'Qualidade da modelagem: 3 decisões centrais')
dec = [
    ('Habilitação ≠ Alocação',
     'PROFESSOR_DISCIPLINA registra quem PODE ministrar (capacidade, vale sempre). '
     'ALOCACAO registra quem MINISTRA, onde e quando (decisão do semestre). '
     'Fatos diferentes → tabelas diferentes.'),
    ('Herança mapeada com 1:1',
     'PROFESSOR e ALUNO têm FK com UNIQUE para USUARIO: dados comuns (nome, login, '
     'senha) existem uma única vez. Sem redundância, sem anomalias de atualização.'),
    ('Restrições viram estrutura',
     '"2 horários por dia" virou UQ(id_dia, slot) com CHECK slot IN (1,2). '
     '"1 professor por disciplina" virou UQ(id_disciplina, id_semestre) em ALOCACAO. '
     'A regra de negócio está no esquema, não só no código.'),
]
xx = Inches(0.55)
w = Inches(4.05)
for i, (t, d) in enumerate(dec):
    card(s, xx, Inches(1.95), w, Inches(4.4), fill=LIGHT)
    tx(s, xx + Inches(0.25), Inches(2.25), w - Inches(0.5), Inches(0.9),
       t, size=17, bold=True, color=NAVY)
    tx(s, xx + Inches(0.25), Inches(3.15), w - Inches(0.5), Inches(3.0),
       d, size=13.5, color=DARK, line_sp=1.15)
    xx += w + Inches(0.18)

# ================= SLIDE 6 - NORMALIZACAO =================
s = slide()
titulo(s, 'Modelo lógico: 13 tabelas na 3ª Forma Normal')
fn = [
    ('1FN', 'Atributos atômicos',
     'Nenhum campo multivalorado ou composto. Dias e horários são linhas de tabelas '
     'próprias (dia_semana, horario), não listas em texto.'),
    ('2FN', 'Sem dependência parcial',
     'A única chave composta é a de professor_disciplina — que não possui atributos '
     'fora da chave. Nada depende de "metade" de uma chave.'),
    ('3FN', 'Sem dependência transitiva',
     'Nome do professor só em usuario; período da disciplina só em disciplina; '
     'capacidade só em sala. Nada descritivo é repetido em alocacao.'),
]
xx = Inches(0.55)
w = Inches(4.05)
for sigla, t, d in fn:
    card(s, xx, Inches(1.9), w, Inches(3.4), fill=LIGHT)
    tx(s, xx + Inches(0.25), Inches(2.15), w - Inches(0.5), Inches(0.8),
       sigla, size=30, bold=True, color=NAVY, font='Georgia')
    tx(s, xx + Inches(0.25), Inches(2.95), w - Inches(0.5), Inches(0.5),
       t, size=14.5, bold=True, color=DARK)
    tx(s, xx + Inches(0.25), Inches(3.5), w - Inches(0.5), Inches(1.7),
       d, size=12.5, color=GRAY, line_sp=1.12)
    xx += w + Inches(0.18)
tx(s, Inches(0.6), Inches(5.65), Inches(12.1), Inches(1.2),
   [('usuario · curso · periodo · disciplina · professor · professor_disciplina · semestre_letivo',
     {'font': 'Consolas', 'size': 13, 'sp_after': 2, 'align': PP_ALIGN.CENTER}),
    ('dia_semana · horario · sala · aluno · alocacao · matricula',
     {'font': 'Consolas', 'size': 13, 'align': PP_ALIGN.CENTER})], color=NAVY)

# ================= SLIDE 7 - DDL / ALOCACAO =================
s = slide()
titulo(s, 'Modelo físico: ALOCACAO concentra as regras',
       'Cada UNIQUE materializa uma regra do enunciado — o SGBD bloqueia o conflito, não o programador')
ddl = """CREATE TABLE alocacao (
  id_alocacao    INT IDENTITY(1,1) NOT NULL,
  id_disciplina  INT NOT NULL,  -- FK disciplina
  id_professor   INT NOT NULL,  -- FK professor
  id_horario     INT NOT NULL,  -- FK horario
  id_sala        INT NOT NULL,  -- FK sala
  id_semestre    INT NOT NULL,  -- FK semestre_letivo
  ...
  CONSTRAINT uq_aloc_disciplina
    UNIQUE (id_disciplina, id_semestre),
  CONSTRAINT uq_aloc_professor
    UNIQUE (id_professor, id_horario, id_semestre),
  CONSTRAINT uq_aloc_sala
    UNIQUE (id_sala, id_horario, id_semestre)
);"""
codebox(s, Inches(0.55), Inches(1.95), Inches(7.6), Inches(4.9), ddl, size=12.5)
notas = [
    ('UQ(disciplina, semestre)', 'uma disciplina tem um único professor — uma só alocação por semestre'),
    ('UQ(professor, horário, semestre)', 'professor não pode estar em dois lugares ao mesmo tempo'),
    ('UQ(sala, horário, semestre)', 'sala não recebe duas turmas simultâneas'),
]
yy = Inches(2.1)
for t, d in notas:
    card(s, Inches(8.4), yy, Inches(4.4), Inches(1.25), fill=ICE)
    tx(s, Inches(8.62), yy + Inches(0.14), Inches(3.96), Inches(1.0),
       [(t, {'bold': True, 'size': 13, 'color': NAVY, 'font': 'Consolas', 'sp_after': 2}),
        (d, {'size': 12, 'color': DARK})])
    yy += Inches(1.42)

# ================= SLIDE 8 - TRIGGERS =================
s = slide()
titulo(s, 'Regra que constraint não alcança → trigger',
       '"Mesmo período no mesmo horário" depende de OUTRA tabela (o período está em disciplina)')
trg = """CREATE TRIGGER trg_valida_periodo_horario ON alocacao
AFTER INSERT, UPDATE AS
BEGIN
    IF EXISTS (
        SELECT 1
        FROM inserted i
        JOIN disciplina di ON di.id_disciplina = i.id_disciplina
        JOIN alocacao a    ON a.id_horario  = i.id_horario
                          AND a.id_semestre = i.id_semestre
                          AND a.id_alocacao <> i.id_alocacao
        JOIN disciplina da ON da.id_disciplina = a.id_disciplina
        WHERE da.id_periodo = di.id_periodo )
    BEGIN
        RAISERROR ('Conflito: ja existe disciplina do mesmo periodo
                    alocada neste horario.', 16, 1);
        ROLLBACK TRANSACTION;
    END
END;"""
codebox(s, Inches(0.55), Inches(1.95), Inches(7.9), Inches(5.0), trg, size=12)
tx(s, Inches(8.75), Inches(2.2), Inches(4.0), Inches(4.6),
   [('Defesa em camadas', {'bold': True, 'size': 17, 'color': NAVY, 'sp_after': 8}),
    ('A aplicação valida e dá mensagens amigáveis ao usuário.', {'sp_after': 8}),
    ('O SGBD garante a integridade mesmo em acesso direto pelo SSMS: '
     'constraints + 2 triggers.', {'sp_after': 8}),
    ('2º trigger: trg_valida_habilitacao — só professor habilitado pode ser alocado.',
     {'color': GRAY, 'size': 13})], size=14.5, line_sp=1.15)

# ================= SLIDE 9 - CONSULTAS PANORAMA =================
s = slide()
titulo(s, 'Consultas SQL: 12 consultas + 2 visões',
       'Todas com no mínimo 3 tabelas · todas executadas e validadas com dados de teste')
grupos = [
    ('2', 'Junções', 'grade completa (8 tabelas) e disciplinas por aluno (7 tabelas)'),
    ('3', 'Conjuntos', 'UNION, INTERSECT e EXCEPT — vínculos, habilitados alocados, capacidade ociosa'),
    ('4', 'Agregações', 'SUM, COUNT, MAX, MIN, AVG — duas com GROUP BY + HAVING'),
    ('3', 'Operadores', "LIKE '%Dados%', BETWEEN datas de admissão, IN ('SEG','SEX')"),
    ('2', 'Visões', 'vw_grade_horarios (9 tabelas) e vw_carga_professores (5 tabelas)'),
]
xx = Inches(0.55)
w = Inches(2.35)
for n, t, d in grupos:
    card(s, xx, Inches(2.1), w, Inches(3.6), fill=LIGHT)
    tx(s, xx + Inches(0.1), Inches(2.35), w - Inches(0.2), Inches(0.9),
       n, size=44, bold=True, color=NAVY, font='Georgia', align=PP_ALIGN.CENTER)
    tx(s, xx + Inches(0.1), Inches(3.35), w - Inches(0.2), Inches(0.5),
       t, size=15, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    tx(s, xx + Inches(0.18), Inches(3.9), w - Inches(0.36), Inches(1.7),
       d, size=11.5, color=GRAY, align=PP_ALIGN.CENTER, line_sp=1.1)
    xx += w + Inches(0.16)
tx(s, Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.5),
   'Resultado dos testes: 14/14 consultas executadas com sucesso, todas com resultados não vazios.',
   size=14, bold=True, color=GREEN, align=PP_ALIGN.CENTER)

# ================= SLIDE 10 - CONSULTA DESTAQUE =================
s = slide()
titulo(s, 'Exemplo: EXCEPT revela a capacidade ociosa',
       'C3 — pares (professor, disciplina) habilitados mas não alocados no semestre')
c3 = """SELECT u.nome AS professor, d.codigo AS disciplina
FROM usuario u
JOIN professor pr            ON pr.id_usuario   = u.id_usuario
JOIN professor_disciplina pd ON pd.id_professor = pr.id_professor
JOIN disciplina d            ON d.id_disciplina = pd.id_disciplina

EXCEPT

SELECT u.nome, d.codigo
FROM usuario u
JOIN professor pr ON pr.id_usuario   = u.id_usuario
JOIN alocacao a   ON a.id_professor  = pr.id_professor
JOIN disciplina d ON d.id_disciplina = a.id_disciplina;"""
codebox(s, Inches(0.55), Inches(2.0), Inches(7.4), Inches(4.3), c3, size=12)
tx(s, Inches(8.3), Inches(2.2), Inches(4.4), Inches(4.4),
   [('Como funciona', {'bold': True, 'size': 16, 'color': NAVY, 'sp_after': 6}),
    ('1º bloco: todas as habilitações (quem PODE ministrar o quê).', {'sp_after': 6}),
    ('2º bloco: alocações efetivas (quem ministra).', {'sp_after': 6}),
    ('EXCEPT devolve a diferença: habilitações não usadas → 5 pares no teste. '
     'Útil para substituições e planejamento.', {'sp_after': 10}),
    ('5 tabelas envolvidas · comparação por tupla inteira', {'size': 12.5, 'color': GRAY})],
   size=14, line_sp=1.15)

# ================= SLIDE 11 - ALGORITMO =================
s = slide()
titulo(s, 'Algoritmo: alocação como coloração de grafos',
       'Disciplinas = vértices · conflito (mesmo período OU mesmo professor) = aresta · horário = cor')
pseudo = """// Fase 1 - professores (guloso por restritividade)
ordenar disciplinas por nº de habilitados (crescente)
para cada disciplina: prof <- habilitado com menor carga

// Fase 2 - grafo de conflitos
aresta {d1,d2} se periodo(d1)=periodo(d2)
             ou prof(d1)=prof(d2)

// Fase 3 - coloracao (Welsh-Powell)
ordenar vertices por grau (decrescente)
para cada vertice: menor cor livre entre vizinhos

// Fase 4 - salas
em cada slot, uma sala distinta por disciplina"""
codebox(s, Inches(0.55), Inches(2.0), Inches(6.9), Inches(4.6), pseudo, size=13)
tx(s, Inches(7.85), Inches(2.15), Inches(4.9), Inches(4.6),
   [('Por que guloso?', {'bold': True, 'size': 16, 'color': NAVY, 'sp_after': 5}),
    ('Coloração mínima é NP-difícil; Welsh-Powell roda em O(V²) e produz '
     'ótimas soluções na prática.', {'sp_after': 10}),
    ('Por que atribuir professor antes?', {'bold': True, 'size': 16, 'color': NAVY, 'sp_after': 5}),
    ('As arestas docentes dependem de quem ministra. Balancear a carga '
     'reduz arestas → favorece o paralelismo.', {'sp_after': 10}),
    ('Coloração válida = grade sem conflitos. Menos cores = mais disciplinas em paralelo.',
     {'size': 13.5, 'color': GRAY})], size=14, line_sp=1.15)

# ================= SLIDE 12 - RESULTADOS =================
s = slide()
titulo(s, 'Resultados: ótimo teórico atingido')
stat(s, Inches(0.7), Inches(2.0), Inches(2.9), '20', 'disciplinas alocadas\n(5 períodos × 4)')
stat(s, Inches(3.8), Inches(2.0), Inches(2.9), '5/10', 'horários usados\n(limite inferior = 5)')
stat(s, Inches(6.9), Inches(2.0), Inches(2.9), '4,0', 'disciplinas em paralelo\n(média por horário)')
stat(s, Inches(10.0), Inches(2.0), Inches(2.9), '0', 'restrições violadas\n(verificação automática)')
grade_txt = """SEG 19:00  P1 IADS     P2 BD1      P3 ENG1     P5 NUV1
SEG 20:55  P1 ALG1     P2 ALG2     P3 BD2      P4 MOB1     P5 TI1
TER 19:00  P1 LOG1     P2 POO1     P3 WEB1     P4 GPR1     P5 IA1
TER 20:55  P1 MTD1     P3 SO1      P4 RED1     P5 TOP1
QUA 19:00  P2 ARC1     P4 SEG1"""
codebox(s, Inches(1.5), Inches(4.35), Inches(10.3), Inches(2.0), grade_txt, size=14)
tx(s, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.5),
   'O grafo contém cliques de tamanho 5 → pelo menos 5 horários são necessários. O algoritmo usou exatamente 5.',
   size=13.5, color=GRAY, align=PP_ALIGN.CENTER)

# ================= SLIDE 13 - APLICACAO =================
s = slide()
titulo(s, 'Aplicação web: 4 funcionalidades sobre o banco',
       'Python + Flask + pyodbc → SQL Server · consultas 100% parametrizadas (sem SQL Injection)')
funcs = [
    ('F1 · Cadastro de usuários', 'validação no código: e-mail, senha mínima, login único — antes do INSERT'),
    ('F2 · Autenticação', 'login/senha verificados no banco; sessão Flask protege todas as rotas'),
    ('F3 · Habilitações', 'manutenção do N:M professor × disciplina, com checagem de duplicidade'),
    ('F4 · Geração da grade', 'algoritmo aciona DELETE + 20 INSERTs em transação atômica; exibição via vw_grade_horarios'),
]
yy = Inches(2.0)
for i, (t, d) in enumerate(funcs):
    card(s, Inches(0.7), yy, Inches(11.9), Inches(1.08), fill=LIGHT if i % 2 == 0 else ICE)
    tx(s, Inches(1.0), yy + Inches(0.16), Inches(3.6), Inches(0.8),
       t, size=15.5, bold=True, color=NAVY)
    tx(s, Inches(4.8), yy + Inches(0.16), Inches(7.6), Inches(0.8),
       d, size=13.5, color=DARK, anchor=MSO_ANCHOR.TOP)
    yy += Inches(1.22)

# ================= SLIDE 14 - CONCLUSAO =================
s = slide(NAVY)
tx(s, Inches(1), Inches(0.9), Inches(11.3), Inches(0.9),
   'Conclusão', size=38, bold=True, color=WHITE, font='Georgia', align=PP_ALIGN.CENTER)
conc = [
    ('Modelar formalmente simplifica', 'enxergar a grade como coloração de grafos transformou uma exigência complexa em um algoritmo simples e verificável.'),
    ('Regras de negócio no esquema', 'UNIQUEs e triggers garantem integridade mesmo fora da aplicação — o banco se defende sozinho.'),
    ('Visões encapsulam complexidade', 'junções de até 9 tabelas viraram consultas de uma linha na aplicação.'),
]
yy = Inches(2.2)
for t, d in conc:
    card(s, Inches(1.6), yy, Inches(10.1), Inches(1.25), fill=RGBColor(0x2A, 0x35, 0x7A))
    tx(s, Inches(1.95), yy + Inches(0.15), Inches(3.6), Inches(1.0),
       t, size=15.5, bold=True, color=ICE)
    tx(s, Inches(5.75), yy + Inches(0.15), Inches(5.7), Inches(1.0),
       d, size=13, color=WHITE, line_sp=1.1)
    yy += Inches(1.45)
tx(s, Inches(1), Inches(6.7), Inches(11.3), Inches(0.5),
   'Trabalhos futuros: múltiplos cursos e turnos · preferências de professores · comparação com métodos exatos',
   size=13, color=ICE, align=PP_ALIGN.CENTER)

saida = os.path.join(DOC, 'Apresentacao_Trabalho_BD.pptx')
prs.save(saida)
print('Apresentacao gerada:', saida)
